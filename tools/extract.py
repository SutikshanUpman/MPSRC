"""Stdlib + python-docx extractor. Writes one UTF-8 .txt per doc into OUT dir.
Usage: python -I extract.py <docs_dir> <out_dir>
"""
import re, sys, zlib, json, os
from pathlib import Path

docs, out = Path(sys.argv[1]), Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)

def pdf_text(data: bytes) -> str:
    chunks = []
    for m in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", data, re.S):
        raw = m.group(1)
        try:
            s = zlib.decompress(raw)
        except Exception:
            s = raw
        if b"BT" not in s:
            continue
        lines = []
        for bt in re.finditer(rb"BT(.*?)ET", s, re.S):
            parts = []
            for tok in re.finditer(rb"\((?:\\.|[^\\)])*\)\s*Tj|\[(.*?)\]\s*TJ|T\*|Td|TD|'", bt.group(1), re.S):
                t = tok.group(0)
                if t.endswith(b"Tj"):
                    parts.append(unesc(t[t.index(b"(")+1:t.rindex(b")")]))
                elif t.endswith(b"TJ"):
                    parts.append(b"".join(unesc(x[1:-1]) for x in re.findall(rb"\((?:\\.|[^\\)])*\)", tok.group(1))))
                else:
                    parts.append(b"\n")
            lines.append(b"".join(parts))
        chunks.append(b"\n".join(lines))
    return b"\n".join(chunks).decode("latin-1")

def unesc(b: bytes) -> bytes:
    b = re.sub(rb"\\([0-7]{1,3})", lambda m: bytes([int(m.group(1), 8) & 255]), b)
    return (b.replace(rb"\(", b"(").replace(rb"\)", b")").replace(rb"\n", b"\n")
             .replace(rb"\r", b"").replace(rb"\t", b"\t").replace(b"\\\\", b"\\"))

def docx_text(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    paras = [x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            paras.append(" | ".join(c.text for c in r.cells))
    return "\n".join(paras)

def txt_text(b: bytes) -> str:
    # header is ascii; body may switch to UTF-16 (mid-file BOM)
    i = b.find(b"\xff\xfe")
    if i > 0:
        return b[:i].decode("utf-8", "replace") + b[i+2:].decode("utf-16-le", "replace") + "\n[[MIXED_ENCODING: utf-16le body at byte %d]]" % i
    return b.decode("utf-8", "replace")

report = {}
for f in sorted(docs.iterdir()):
    b = f.read_bytes()
    try:
        if f.suffix == ".pdf":
            t = pdf_text(b)
        elif f.suffix == ".docx":
            t = docx_text(f)
        else:
            t = txt_text(b)
        status = "ok" if t.strip() else "EMPTY"
    except Exception as e:
        t, status = "", f"ERROR {type(e).__name__}: {e}"
    (out / (f.stem + ".txt")).write_text(t, encoding="utf-8")
    report[f.name] = {"bytes": len(b), "chars": len(t), "status": status}
print(json.dumps(report, indent=0))
